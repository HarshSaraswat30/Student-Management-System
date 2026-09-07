from openpyxl import Workbook, load_workbook
import os
import sys
import msvcrt as mv
FileName="newfile.xlsx"
if not os.path.exists(FileName):
    #state = True
    wb = Workbook()
    ws1 = wb.active

    ws1.title = "student"

    ws1['A1'] = "Student ID"
    ws1['B1'] = "Student Name"
    ws1['C1'] = "Student Class"
    ws1['D1'] = "Student City"
    ws1['E1'] = "Student Phone No"
    wb.save(FileName)
    #print()

def insertData():
    os.system("cls")
    wb = load_workbook(FileName)
    ws1 = wb.active
    data = []
    tid = input("Enter ID -- ")
    insstat = False
    for row in ws1.iter_rows(min_row=2, values_only=True):
        if tid == row[0]:
            insstat = True
            print("ID Already Exists.. press enter to continue.. ")
            mv.getch()
            break
    if insstat == True:
        insertData()
    data.append(tid)
    data.append(input("Enter Name -- "))
    data.append(input("Enter Class -- "))
    data.append(input("Enter City -- "))
    data.append(input("Enter Phone NO -- "))
    # data = ["101","Harsh Saraswat","BCA","Ajmer","9852158478"]
    ws1.append(data)
    wb.save(FileName)
    print("\nRecord Inserted Successfully.")
    input()

def showData():
    os.system("cls")
    wb = load_workbook(FileName)
    ws1 = wb.active
    print("Showing the list module called....  \n")
    print("=" * 95)
    print(f"|{'ID':<10}|{'NAME':<25}|{'CLASS':<15}|{'CITY':<15}|{'PHONE':<20}|")
    print("=" * 95)

    for row in ws1.iter_rows(min_row=2, values_only=True):
        print(f"|{str(row[0]):<10}|{str(row[1]):<25}|{str(row[2]):<15}|{str(row[3]):<15}|{str(row[4]):<20}|")

    print("-" * 95)

    input("\nPress Enter...")

def DeleteData():
    os.system("cls")

    wb = load_workbook(FileName)
    ws1 = wb.active

    empid = input("Enter Student ID to Delete : ").strip()

    found = False

    for row in range(2, ws1.max_row + 1):

        if str(ws1.cell(row=row, column=1).value).strip() == empid:

            found = True

            print("-" * 50)
            print("Student ID    :", ws1.cell(row, 1).value)
            print("Student Name  :", ws1.cell(row, 2).value)
            print("Student Class :", ws1.cell(row, 3).value)
            print("Student City  :", ws1.cell(row, 4).value)
            print("Student Phone :", ws1.cell(row, 5).value)
            print("-" * 50)

            print("1. Temporary Delete")
            print("2. Permanent Delete")
            print("3. Skip")

            ch = mv.getch().decode()

            if ch == '1':

                data = [
                    ws1.cell(row, 1).value,
                    ws1.cell(row, 2).value,
                    ws1.cell(row, 3).value,
                    ws1.cell(row, 4).value,
                    ws1.cell(row, 5).value
                ]

                wb1 = load_workbook("recycle.xlsx")
                ws2 = wb1.active
                ws2.append(data)
                wb1.save("recycle.xlsx")

                ws1.delete_rows(row)
                wb.save(FileName)

                print("\nRecord Moved to Recycle Bin Successfully.")

            elif ch == '2':

                ws1.delete_rows(row)
                wb.save(FileName)

                print("\nRecord Permanently Deleted.")

            elif ch == '3':

                print("\nDelete Operation Skipped.")

            break

    if not found:
        print("\nStudent ID Not Found.")

    print("\nDo you want to continue?")
    print("Press Y = Delete Again")
    print("Press M = Main Menu")
    print("Any Key = Exit")

    c = mv.getch().decode().upper()

    if c == 'Y':
        DeleteData()
    elif c == 'M':
        MainModule()
    else:
        exit()

def UpdateData():
    os.system("cls")

    wb = load_workbook(FileName)
    ws1 = wb.active

    up = input("Enter Student ID : ").strip()

    found = False

    for i in range(2, ws1.max_row + 1):

        if str(ws1.cell(i, 1).value).strip() == up:

            found = True

            print("\nOld Record")
            print("-" * 40)
            print("ID    :", ws1.cell(i, 1).value)
            print("Name  :", ws1.cell(i, 2).value)
            print("Class :", ws1.cell(i, 3).value)
            print("City  :", ws1.cell(i, 4).value)
            print("Phone :", ws1.cell(i, 5).value)
            print("-" * 40)

            print("\n1. Update Name")
            print("2. Update Class")
            print("3. Update City")
            print("4. Update Phone")
            print("5. Update All")

            ch = input("Enter Choice : ")

            if ch == "1":
                ws1.cell(i, 2).value = input("Enter New Name : ")

            elif ch == "2":
                ws1.cell(i, 3).value = input("Enter New Class : ")

            elif ch == "3":
                ws1.cell(i, 4).value = input("Enter New City : ")

            elif ch == "4":
                ws1.cell(i, 5).value = input("Enter New Phone : ")

            elif ch == "5":
                ws1.cell(i, 2).value = input("Enter New Name : ")
                ws1.cell(i, 3).value = input("Enter New Class : ")
                ws1.cell(i, 4).value = input("Enter New City : ")
                ws1.cell(i, 5).value = input("Enter New Phone : ")

            else:
                print("Invalid Choice")
                wb.close()
                return
            i=input("\nAre you sure you want to update this record? (Y/N): ").strip().upper()
            if i =="Y":
                wb.save(FileName)
                print("\nRecord Updated Successfully.")
                break

    if not found:
        print("\nRecord Not Found.")

    wb.close()

    # input("\nPress Enter...")


def showRecData():

    os.system("cls")
    wb = load_workbook("recycle.xlsx")
    ws1 = wb.active
    print("Showing the list module called....  \n")
    print("=" * 95)
    print(f"|{'ID':<10}|{'NAME':<25}|{'CLASS':<15}|{'CITY':<15}|{'PHONE':<20}|")
    print("=" * 95)

    for row in ws1.iter_rows(min_row=2, values_only=True):
        print(f"|{str(row[0]):<10}|{str(row[1]):<25}|{str(row[2]):<15}|{str(row[3]):<15}|{str(row[4]):<20}|")

    print("-" * 95)

    input("\nPress Enter...")

def recycleRecover():

    os.system("cls")

    print("===== Recycle Bin =====")
    print("1. Recover One Data")
    print("2. Recover All Data")
    print("3. Delete One Data")
    print("4. Delete All Data")
    print("5. Back")
    print("=" * 25)

    ch = input("Enter Choice : ")

    if ch == "5":
        return

    main_wb = load_workbook("newfile.xlsx")
    main_ws = main_wb.active

    recycle_wb = load_workbook("recycle.xlsx")
    recycle_ws = recycle_wb.active

    if recycle_ws.max_row <= 1:
        print("\nRecycle Bin Empty.")
        input("\nPress Enter...")
        return

    # Show Recycle Data
    print("\nRecycle Bin Data")
    print("-" * 80)
    print(f"{'ID':<10}{'NAME':<20}{'CLASS':<15}{'CITY':<15}{'PHONE':<15}")
    print("-" * 80)

    for row in recycle_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            print(f"{row[0]:<10}{row[1]:<20}{row[2]:<15}{row[3]:<15}{row[4]:<15}")

    print("-" * 80)

    # ---------------- RECOVER ONE ----------------

    if ch == "1":

        sid = input("\nEnter Student ID : ")
        found = False

        for i in range(2, recycle_ws.max_row + 1):

            if str(recycle_ws.cell(i, 1).value) == sid:

                found = True

                ans = input("Recover this record (Y/N): ").upper()

                if ans == "Y":

                    data = [
                        recycle_ws.cell(i, 1).value,
                        recycle_ws.cell(i, 2).value,
                        recycle_ws.cell(i, 3).value,
                        recycle_ws.cell(i, 4).value,
                        recycle_ws.cell(i, 5).value
                    ]

                    main_ws.append(data)

                    recycle_ws.delete_rows(i)

                    main_wb.save("newfile.xlsx")
                    recycle_wb.save("recycle.xlsx")

                    print("\nRecord Recovered Successfully.")

                else:
                    print("\nRecover Cancelled.")

                break

        if not found:
            print("\nStudent ID Not Found.")

    # ---------------- RECOVER ALL ----------------

    elif ch == "2":

        ans = input("\nRecover All Records (Y/N): ").upper()

        if ans == "Y":

            count = 0

            while recycle_ws.max_row > 1:

                data = [
                    recycle_ws.cell(2, 1).value,
                    recycle_ws.cell(2, 2).value,
                    recycle_ws.cell(2, 3).value,
                    recycle_ws.cell(2, 4).value,
                    recycle_ws.cell(2, 5).value
                ]

                main_ws.append(data)

                recycle_ws.delete_rows(2)

                count += 1

            main_wb.save("newfile.xlsx")
            recycle_wb.save("recycle.xlsx")

            print(f"\n{count} Records Recovered Successfully.")

        else:
            print("\nRecover Cancelled.")

    # ---------------- DELETE ONE ----------------

    elif ch == "3":

        sid = input("\nEnter Student ID : ")
        found = False

        for i in range(2, recycle_ws.max_row + 1):

            if str(recycle_ws.cell(i, 1).value) == sid:

                found = True

                ans = input("Delete Permanently (Y/N): ").upper()

                if ans == "Y":

                    recycle_ws.delete_rows(i)

                    recycle_wb.save("recycle.xlsx")

                    print("\nRecord Deleted Permanently.")

                else:
                    print("\nDelete Cancelled.")

                break

        if not found:
            print("\nStudent ID Not Found.")

    # ---------------- DELETE ALL ----------------

    elif ch == "4":

        ans = input(
            "\nDelete ALL Records Permanently (Y/N): "
        ).upper()

        if ans == "Y":

            recycle_ws.delete_rows(
                2,
                recycle_ws.max_row - 1
            )

            recycle_wb.save("recycle.xlsx")

            print("\nAll Records Deleted Permanently.")

        else:
            print("\nDelete Cancelled.")

    else:
        print("\nInvalid Choice.")

    input("\nPress Enter...")


def MainModule(): 

    while True:  
        os.system("cls")
   
        print("===== Student System =====")
        print("1. Insert Student")
        print("2. Show   Student")
        print("3. Delete Student")
        print("4. Udapte Student")
        print("5. Show Recycle Student")
        print("6  Recover Student ")
        print("7. Exit")

        ch = mv.getch().decode()

        if ch == "1":
            insertData()
        elif ch == "2":
            showData()
        elif ch == "3":
            DeleteData()
        elif ch == "4":
            UpdateData()
        elif ch == "5":
            showRecData()
        elif ch == "6":
            recycleRecover()
        elif ch =="7":
            print("Program Closed...")
            sys.exit()
        else:
            print("Invalid Choice...")
MainModule()

